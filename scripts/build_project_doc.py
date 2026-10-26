from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path('Expense_Tracker_Mainframe_Project_Explanation.docx')
NAVY, TEAL, GRAY, BORDER = '17365D', '087E81', 'F3F5F7', 'D9D9D9'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); edges = tcPr.first_child_found_in('w:tcBorders')
    if edges is None: edges = OxmlElement('w:tcBorders'); tcPr.append(edges)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = edges.find(qn('w:'+edge))
        if el is None: el = OxmlElement('w:'+edge); edges.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'6'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),BORDER)

def cell_text(cell, text, bold=False, color=None, size=9.2):
    cell.text = ''; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(cell)

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers): cell_text(t.rows[0].cells[i], h, True, 'FFFFFF'); shade(t.rows[0].cells[i], NAVY)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value)
            if ri % 2: shade(cells[i], GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def h(doc, text, level=1):
    p = doc.add_heading(text, level=level); p.paragraph_format.keep_with_next = True; return p

def bullet(doc, text): doc.add_paragraph(text, style='List Bullet')

def code(doc, text):
    p = doc.add_paragraph(style='CodeBlock'); p.add_run(text); return p

doc = Document(); s = doc.sections[0]
s.top_margin = Inches(.7); s.bottom_margin = Inches(.7); s.left_margin = Inches(.8); s.right_margin = Inches(.8)
normal = doc.styles['Normal']; normal.font.name = 'Aptos'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos'); normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(31,41,51); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.08
for name, size, color in [('Title',28,NAVY),('Heading 1',18,NAVY),('Heading 2',13,TEAL),('Heading 3',11,NAVY)]:
    st=doc.styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(6)
cb = doc.styles.add_style('CodeBlock', WD_STYLE_TYPE.PARAGRAPH); cb.font.name='Consolas'; cb.font.size=Pt(8.5); cb.paragraph_format.left_indent=Inches(.2); cb.paragraph_format.space_after=Pt(7)

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Expense Tracker Mainframe Project')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Complete Architecture and Module Explanation'); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=RGBColor.from_string(TEAL)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(35); p.add_run('A browser-based expense workflow integrated with FastAPI, Zowe CLI, z/OS JCL, Enterprise COBOL, and VSAM').font.size=Pt(12)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(45); p.add_run('Prepared for project explanation, demonstration, and five-member presentation').italic=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(75); p.add_run('Project namespace: Z86073.EXPTRK.*\nVerified compiler: Enterprise COBOL 6.5 with IGYWCL\nVerified live path: UI to FastAPI to Zowe to z/OS').font.size=Pt(11)
doc.add_page_break()

h(doc,'Document Purpose'); doc.add_paragraph('This document explains how the Expense Tracker works as one complete system. It separates local application responsibilities from z/OS responsibilities, describes where data is stored, explains the approval and finance batch workflow, and assigns the project modules across five team members. The main conclusion is that the mainframe VSAM expense master is the authoritative store for posted expenses; the browser and FastAPI layer coordinate requests, permissions, job monitoring, and verified results.')
h(doc,'Project Summary'); add_table(doc,['Area','Implementation','Responsibility'],[
('User interface','Node-served HTML and JavaScript','Login, forms, approvals, batch controls, reports'),('Application bridge','FastAPI','Validation, request state, Zowe orchestration, result verification'),('Mainframe access','Zowe CLI and z/OSMF','Upload members, submit JCL, retrieve JES output and datasets'),('Business programs','Enterprise COBOL 6.5','Post, read, and report expenses'),('Authoritative storage','VSAM KSDS','Employee master and posted expense master'),('Operations','JCL and REXX','Allocation, compilation, batch execution, monitoring')])
h(doc,'End to End Architecture'); doc.add_paragraph('The browser never receives Zowe credentials and never executes JCL directly. The Node server authenticates the user and exposes an allowlisted same-origin proxy. FastAPI performs controlled server-side calls to Zowe CLI. z/OS then executes the approved JCL and COBOL programs.')
code(doc,'Employee browser\n    | login, submit, approve, batch, reports\n    v\nNode.js UI server and session layer\n    | allowlisted /mainframe proxy\n    v\nFastAPI integration service\n    | Zowe CLI subprocesses over z/OSMF\n    v\nz/OS JCL jobs\n    +--> EXPPOST --> Employee KSDS + Expense KSDS + RESULT\n    +--> EXPREAD --> Expense KSDS --> READOUT\n    +--> EXPRPT  --> Expense KSDS --> REPORT\n    +--> EXPMON  --> bounded operational checks')

h(doc,'Local Application Components'); add_table(doc,['Component','Important files','What it does'],[
('Browser UI','public/index.html, public/live-workflow.js','Role-based pages, forms, approvals, batches, reports, status messages.'),('Node server','server.js','Serves UI, manages sessions, validates permissions, proxies live paths.'),('FastAPI','backend/app/main.py','Health, expenses, approvals, batches, jobs, and reports endpoints.'),('Zowe adapter','backend/app/services/expense_service.py','Allocates input, uploads records, submits JCL, polls jobs, verifies results.'),('Runtime tracking','backend/runtime/jobs.json and batches.json','Local workflow metadata; not a replacement for VSAM.'),('Demo mode','data/users.json and data/workflow-demo.json','Optional local demonstration mode, separate from live mode.')])
h(doc,'Live API Workflow'); add_table(doc,['Endpoint','Purpose','Main users'],[
('GET /api/health','Application health and configuration','All'),('GET /api/mainframe/health','Mainframe integration check','Finance/Admin'),('POST /api/expenses','Creates pending approval expense','Employee'),('GET /api/expenses','Loads report-backed records','Manager/Finance'),('POST /api/expenses/{id}/approve','Approves expense','Manager'),('POST /api/expenses/{id}/reject','Rejects expense','Manager'),('POST /api/batches','Processes approved expenses','Finance/Admin'),('GET /api/batches','Shows batch status and counts','Finance/Admin'),('GET /api/jobs/{request_id}','Shows job state','Authorized users'),('GET /api/reports/mainframe','Reads verified mainframe report','Manager/Finance/Admin')])

h(doc,'Business Workflow'); h(doc,'Employee Submission',2); doc.add_paragraph('An employee enters a unique expense ID, employee ID, date, category, amount, and description. The API validates the request and returns pending_approval. The record is not posted to VSAM at this stage because employees are not allowed to directly update the expense master.')
h(doc,'Manager Approval',2); doc.add_paragraph('A manager opens Approvals and chooses Approve or Reject. Approval is an application-level authorization event. It authorizes finance to include the expense in a controlled batch; it is not itself a mainframe posting.')
h(doc,'Finance Batch Processing',2); doc.add_paragraph('Finance starts a batch containing approved records. FastAPI creates a unique transaction input dataset, uploads the fixed-width record, submits EXPPOST, polls the JES job, reads the result dataset, and refreshes the report with EXPRPT. A record is shown as Posted only after the matching result record confirms the mainframe operation.')
h(doc,'Why a Batch Is Used',2)
for x in ['It separates employee intake from finance-controlled posting.','It creates a controlled and auditable processing boundary.','It allows approved records to be processed together and counted.','It keeps JCL submission and mainframe credentials on the server.','It gives finance clear posted, rejected, and failed counts.']: bullet(doc,x)

h(doc,'Mainframe Components'); h(doc,'Employee VSAM KSDS',2); add_table(doc,['Property','Value'],[('Dataset','Z86073.EXPTRK.EMPLOYEE'),('Organization','VSAM KSDS'),('Record format','FB80'),('Key','Employee ID'),('Key offset','0'),('Key length','8 bytes'),('Example','E0001002'),('Status rule','Only active employees are accepted')]); doc.add_paragraph('Employee record layout: ID X(8), NAME X(30), DEPT X(10), STATUS X(1), FILLER X(31). Total length is exactly 80 bytes.')
h(doc,'Expense VSAM KSDS',2); add_table(doc,['Property','Value'],[('Dataset','Z86073.EXPTRK.MASTER'),('Components','Z86073.EXPTRK.MASTER.DATA and .INDEX'),('Organization','VSAM KSDS'),('Record format','FB80'),('Key','Expense ID'),('Key offset','0'),('Key length','12 bytes'),('Authority','Authoritative store for posted expenses')]); doc.add_paragraph('Expense record layout: EXPENSE-ID X(12), EMPLOYEE-ID X(8), DATE X(8), CATEGORY X(4), AMOUNT X(9), DESCRIPTION X(39). Total length is exactly 80 bytes. Amount is integer paise using nine zero-padded digits.')
h(doc,'Sequential Supporting Datasets',2); add_table(doc,['Dataset','Use','Record length'],[('Z86073.EXPTRK.TXN.*','Per-request transaction input','80'),('Z86073.EXPTRK.RESULT','EXPPOST result records','120'),('Z86073.EXPTRK.REQUEST','Expense ID request for EXPREAD','12'),('Z86073.EXPTRK.READOUT','EXPREAD result','120'),('Z86073.EXPTRK.REPORT','EXPRPT output consumed by API','120'),('Z86073.EXPTRK.SEED','Controlled seed data','80')])

h(doc,'COBOL Programs'); add_table(doc,['Program','Input','Mainframe action','Output'],[('EXPPOST','TRANSIN, EMPVSAM, EXPVSAM','Validate employee and expense, prevent duplicates, write accepted expense','RESULT'),('EXPREAD','REQUEST, EXPVSAM','Keyed lookup by 12-byte expense ID','READOUT'),('EXPRPT','EXPVSAM','Read master and calculate report totals','RPTOUT / REPORT')]); doc.add_paragraph('All programs use explicit SELECT, FD, ASSIGN, RECORD KEY, FILE STATUS, and DD-name conventions. This prevents a JCL DD name from pointing to the wrong file and ensures that a failed VSAM operation cannot be confused with a successful post.')

h(doc,'JCL and Operations'); add_table(doc,['Member','Purpose'],[('EXPDEF.jcl','Defines the project expense KSDS without touching coursework clusters.'),('EMPDEF.jcl','Defines the project employee KSDS.'),('EXPLOAD.jcl','Loads controlled seed records after definition.'),('EXPALLOC.jcl','Allocates supporting sequential datasets.'),('EXPPOST.jcl','Runs EXPPOST with employee and expense VSAM DDs.'),('EXPREAD.jcl','Runs keyed expense readback.'),('EXPRPT.jcl','Runs EXPRPT with the RPTOUT DD name.'),('EXPMON.rexx','Performs bounded monitoring and result checks.')]); doc.add_paragraph('Compilation uses Enterprise COBOL 6.5 and IGYWCL. Load modules are placed in Z86073.EXPTRK.LOAD. Zowe CLI submits jobs and retrieves JES output through z/OSMF.')

h(doc,'Verified Mainframe Evidence'); add_table(doc,['Verification','Evidence'],[('Expense master provisioning','JOB04562, CC 0000; KEYS(12 0), RECORDSIZE(80 80).'),('Employee master','JOB05375 defined, JOB05378 loaded, JOB05384 verified.'),('Supporting datasets','JOB05388 allocated.'),('COBOL compilation','EXPPOST JOB02996, EXPREAD JOB02999, EXPRPT JOB03002; all CC 0000.'),('Earlier controlled transaction','JOB05414 posted EX26102501A1; JOB05417 returned FOUND; JOB05423 generated report.'),('Live posting','JOB07939 posted EX261026E05 with READ=1, POSTED=1, REJECTED=0.'),('Live report','JOB07942 generated the report containing the verified expense.'),('Approval and batch','EX261026WF1 moved pending approval to approved to completed; posted=1, rejected=0, failed=0.')])

h(doc,'Five Member Module Division'); doc.add_paragraph('Each member should explain the purpose, important files or datasets, input and output, and one demonstration step.')
add_table(doc,['Member','Assigned module','Topics to explain','Demonstration'],[('Member 1','Frontend and user roles','Login, employee submission, manager approval, finance batch screens, statuses, live versus demo.','Submit an expense and show Pending approval.'),('Member 2','FastAPI backend and security','Endpoints, validation, permissions, asynchronous states, duplicate protection, sanitized errors, server-side credentials.','Show API health, request ID, and permission behavior.'),('Member 3','Mainframe data and VSAM','Employee KSDS, expense KSDS, record layouts, keys, supporting datasets, authoritative storage.','Show dataset names and explain one 80-byte record.'),('Member 4','COBOL, JCL, and Zowe','EXPPOST, EXPREAD, EXPRPT, DD names, IGYWCL, JES jobs, result records.','Show job ID, return code, result, and readback.'),('Member 5','Testing and operations','EXPMON, batch control, evidence, failure handling, dataset safety, future work.','Run manager approval, finance batch, and confirm report.')])
h(doc,'Suggested Presentation Order',2); code(doc,'1. Member 1: User problem and browser workflow\n2. Member 2: API, authentication, authorization, and request states\n3. Member 3: Mainframe datasets and fixed-width/VSAM design\n4. Member 4: COBOL programs, JCL submission, JES, and Zowe\n5. Member 5: Testing evidence, monitoring, report verification, and conclusion')

h(doc,'Manual Demonstration Script'); add_table(doc,['Step','Role','Action','Expected result'],[('1','Employee','Log in and submit a unique expense ID.','Pending approval.'),('2','Manager','Open Approvals and approve the expense.','Approved.'),('3','Finance','Open Batch Operations and run the batch.','Queued/running, then completed.'),('4','Finance','Inspect counts and job status.','Posted count increases; failures are accurate.'),('5','Manager or Finance','Open expenses and reports.','Expense appears as Posted from verified report data.'),('6','Optional technical check','Use EXPREAD/readback for the expense ID.','Mainframe returns FOUND and the exact record.')]); doc.add_paragraph('Use a new unique expense ID for a new test. Never resubmit an already posted ID such as EX26102501A1.')

h(doc,'Important Safety and Limitations')
for x in ['Existing coursework datasets Z86073.VSAM.KSDS and Z86073.VSAMDS are not redefined or used by this project.','Provisioning is separate from runtime posting and is not automatic at API startup.','The API reports unavailable or failed when Zowe, permissions, allocation, JES, VSAM, or result verification fails.','CC 0000 alone is not business proof; the result record, VSAM readback, and report must also be checked.','Local demo mode is not synchronized with live mainframe mode and must be labelled separately.']: bullet(doc,x)
h(doc,'Conclusion'); doc.add_paragraph('The project demonstrates a controlled enterprise workflow rather than a browser-only expense list. Employees submit requests, managers authorize them, finance runs a batch, and the mainframe performs authoritative validation and posting. COBOL owns the business rules at the VSAM boundary; JCL defines program datasets; Zowe provides controlled remote operations; FastAPI verifies actual results; and the browser displays only the state supported by those results.')

for section in doc.sections:
    f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=f.add_run('Expense Tracker Mainframe Project | '); r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
doc.save(OUT); print(OUT.resolve())
